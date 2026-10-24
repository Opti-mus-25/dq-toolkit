import tempfile
import unittest
from pathlib import Path

import pandas as pd

from dqkit.cleaning import clean_table, parse_date
from dqkit.issues import IssueLog
from dqkit.report import DICTIONARY_COLUMNS, build_dictionary, write_xlsx
from dqkit.rules import load_rules
from dqkit.sqlchecks import parse_sql_file
from dqkit.validate import check_computed_totals, check_fields, check_foreign_keys

ROOT = Path(__file__).resolve().parents[1]
RULES = load_rules(ROOT / "rules" / "rules.yaml")


def frame(rows):
    return pd.DataFrame(rows, dtype="string")


class CleaningTests(unittest.TestCase):
    def test_label_variants_are_standardized_and_logged(self):
        log = IssueLog(RULES["categories"])
        raw = frame(
            [
                {"participant_id": "P0001", "birth_date": "1940-01-02", "sex": "female", "education_years": "12", "enrollment_date": "2018-01-01"},
                {"participant_id": "P0002", "birth_date": "1940-01-02", "sex": "M ", "education_years": "12", "enrollment_date": "2018-01-01"},
            ]
        )
        cleaned = clean_table(raw, "participants", RULES["tables"]["participants"], RULES, log)
        self.assertEqual(list(cleaned["sex"]), ["F", "M"])
        self.assertEqual(len(log.rows), 2)

    def test_us_dates_become_iso(self):
        log = IssueLog(RULES["categories"])
        raw = frame(
            [{"participant_id": "P0001", "birth_date": "03/14/1941", "sex": "F", "education_years": "12", "enrollment_date": "2018-01-01"}]
        )
        cleaned = clean_table(raw, "participants", RULES["tables"]["participants"], RULES, log)
        self.assertEqual(cleaned.at[0, "birth_date"], "1941-03-14")

    def test_exact_duplicates_are_removed_once(self):
        log = IssueLog(RULES["categories"])
        row = {"participant_id": "P0001", "birth_date": "1940-01-02", "sex": "F", "education_years": "12", "enrollment_date": "2018-01-01"}
        cleaned = clean_table(frame([row, row]), "participants", RULES["tables"]["participants"], RULES, log)
        self.assertEqual(len(cleaned), 1)
        self.assertEqual(log.rows[0]["category"], "duplicate_removed")

    def test_unrecognised_date_is_left_for_validation(self):
        self.assertIsNone(parse_date("14 March 1941", RULES["date_formats"]))


class ValidationTests(unittest.TestCase):
    def run_fields(self, rows):
        log = IssueLog(RULES["categories"])
        check_fields(frame(rows), "participants", RULES["tables"]["participants"], log)
        return [row["category"] for row in log.rows]

    def base(self, **overrides):
        row = {"participant_id": "P0001", "birth_date": "1940-01-02", "sex": "F", "education_years": "12", "enrollment_date": "2018-01-01"}
        row.update(overrides)
        return row

    def test_valid_row_has_no_issues(self):
        self.assertEqual(self.run_fields([self.base()]), [])

    def test_required_type_range_category_pattern(self):
        self.assertEqual(self.run_fields([self.base(education_years=pd.NA)]), ["missing_required"])
        self.assertEqual(self.run_fields([self.base(education_years="twelve")]), ["invalid_type"])
        self.assertEqual(self.run_fields([self.base(education_years="40")]), ["out_of_range"])
        self.assertEqual(self.run_fields([self.base(sex="X")]), ["invalid_category"])
        self.assertEqual(self.run_fields([self.base(participant_id="P12")]), ["invalid_format"])
        self.assertEqual(self.run_fields([self.base(birth_date="1890-01-01")]), ["out_of_range"])

    def test_total_score_mismatch_and_skip_when_item_missing(self):
        scores = frame(
            [
                {"visit_id": "V00001", "participant_id": "P0001", "orientation": "9", "registration": "3", "attention": "4", "recall": "2", "language": "8", "total_reported": "27"},
                {"visit_id": "V00002", "participant_id": "P0001", "orientation": "9", "registration": "3", "attention": "4", "recall": pd.NA, "language": "8", "total_reported": "99"},
            ]
        )
        log = IssueLog(RULES["categories"])
        check_computed_totals({"cognitive_scores": scores}, RULES, log)
        self.assertEqual([row["record_key"] for row in log.rows], ["V00001"])

    def test_orphan_foreign_key(self):
        tables = {
            "participants": frame([{"participant_id": "P0001"}]),
            "visits": frame(
                [
                    {"visit_id": "V00001", "participant_id": "P0001"},
                    {"visit_id": "V00002", "participant_id": "P9999"},
                ]
            ),
            "cognitive_scores": frame([{"visit_id": "V00001", "participant_id": "P0001"}]),
        }
        log = IssueLog(RULES["categories"])
        check_foreign_keys(tables, RULES, log)
        self.assertEqual([row["record_key"] for row in log.rows], ["V00002"])


class SqlAndReportTests(unittest.TestCase):
    def test_sql_file_parses_into_named_checks(self):
        checks = parse_sql_file(ROOT / "sql" / "cross_checks.sql")
        self.assertEqual(len(checks), 5)
        for check in checks:
            self.assertIn("SELECT", check["sql"])
            self.assertIn(check["category"], RULES["categories"])

    def test_dictionary_uses_redcap_headers_and_unique_names(self):
        dictionary = build_dictionary(RULES)
        self.assertEqual(list(dictionary.columns), DICTIONARY_COLUMNS)
        self.assertTrue(dictionary["Variable / Field Name"].is_unique)
        self.assertEqual(dictionary.iloc[0]["Variable / Field Name"], "participant_id")

    def test_excel_report_has_three_sheets(self):
        from openpyxl import load_workbook

        log = IssueLog(RULES["categories"])
        log.add("out_of_range", "participants", "P0001", "education_years", "40", "Value 40 is above the maximum of 25")
        log.add("label_standardized", "participants", "P0002", "sex", "female", "Standardized", action="auto_fixed")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "report.xlsx"
            write_xlsx(log.frame(), path)
            self.assertEqual(load_workbook(path).sheetnames, ["Summary", "Needs Review", "Auto-fixes"])


if __name__ == "__main__":
    unittest.main()
