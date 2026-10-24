from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEVERITY_ORDER = {"high": 0, "medium": 1, "low": 2}


def split_issues(issues):
    review = issues[issues["action"] == "needs_review"].copy()
    fixes = issues[issues["action"] == "auto_fixed"].copy()
    review["_rank"] = review["severity"].map(SEVERITY_ORDER)
    review = review.sort_values(["_rank", "route_to", "category", "table", "record_key"]).drop(columns="_rank")
    return review.reset_index(drop=True), fixes.reset_index(drop=True)


def summary_frame(issues):
    grouped = (
        issues.groupby(["action", "severity", "category", "route_to"])
        .size()
        .reset_index(name="count")
    )
    grouped["_rank"] = grouped["severity"].map(SEVERITY_ORDER)
    grouped = grouped.sort_values(["action", "_rank", "category"], ascending=[False, True, True])
    return grouped.drop(columns="_rank").reset_index(drop=True)


def write_cleaned(cleaned, rules, out_dir):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    for table, df in cleaned.items():
        df.to_csv(out_dir / rules["tables"][table]["file"], index=False)


def write_issues_csv(issues, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    issues.to_csv(path, index=False)


def _style_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="C5050C")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        longest = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 70)


def write_xlsx(issues, path):
    review, fixes = split_issues(issues)
    summary = summary_frame(issues)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        review.to_excel(writer, sheet_name="Needs Review", index=False)
        fixes.to_excel(writer, sheet_name="Auto-fixes", index=False)
        for sheet in writer.book.worksheets:
            _style_sheet(sheet)


def _table(frame, columns):
    lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join("---" for _ in columns) + " |"]
    for row in frame[columns].itertuples(index=False, name=None):
        lines.append("| " + " | ".join(str(value).replace("|", "/") for value in row) + " |")
    return "\n".join(lines)


def write_markdown(issues, raw_counts, cleaned, path):
    review, fixes = split_issues(issues)
    lines = ["# Data Quality Discrepancy Report", ""]
    lines.append("Synthetic practice data. No real participants.")
    lines.append("")
    lines.append("## Run summary")
    lines.append("")
    counts = pd.DataFrame(
        [
            {"table": table, "rows_read": raw_counts[table], "rows_after_cleaning": len(cleaned[table])}
            for table in raw_counts
        ]
    )
    lines.append(_table(counts, ["table", "rows_read", "rows_after_cleaning"]))
    lines.append("")
    lines.append(f"- Issues needing review: **{len(review)}**")
    lines.append(f"- Fixes applied automatically and logged: **{len(fixes)}**")
    lines.append("")
    lines.append("## Counts by category")
    lines.append("")
    lines.append(_table(summary_frame(issues), ["action", "severity", "category", "route_to", "count"]))
    lines.append("")
    lines.append("## Issues needing review, grouped by who resolves them")
    for team, group in review.groupby("route_to", sort=True):
        lines.append("")
        lines.append(f"### {team} ({len(group)})")
        lines.append("")
        lines.append(_table(group, ["issue_id", "severity", "category", "table", "record_key", "field", "value", "message"]))
    lines.append("")
    lines.append("## Automatic fixes (audit trail)")
    lines.append("")
    fix_counts = fixes.groupby("category").size().reset_index(name="count")
    lines.append(_table(fix_counts, ["category", "count"]))
    lines.append("")
    lines.append("Every automatic fix is listed row by row in `issues_log.csv` and the `Auto-fixes` sheet of the Excel report.")
    lines.append("")
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text("\n".join(lines), encoding="utf-8")


DICTIONARY_COLUMNS = [
    "Variable / Field Name",
    "Form Name",
    "Section Header",
    "Field Type",
    "Field Label",
    "Choices, Calculations, OR Slider Labels",
    "Field Note",
    "Text Validation Type OR Show Slider Number",
    "Text Validation Min",
    "Text Validation Max",
    "Identifier?",
    "Branching Logic (Show field only if...)",
    "Required Field?",
    "Custom Alignment",
    "Question Number (surveys only)",
    "Matrix Group Name",
    "Matrix Ranking?",
    "Field Annotation",
]


def build_dictionary(rules, exclude_tables=("clinical_participants",)):
    rows = []
    seen = set()
    for table, spec in rules["tables"].items():
        if table in exclude_tables:
            continue
        form = spec["form"].lower().replace(" ", "_")
        for field, field_spec in spec["fields"].items():
            if field in seen:
                continue
            seen.add(field)
            kind = field_spec.get("type", "string")
            row = {column: "" for column in DICTIONARY_COLUMNS}
            row["Variable / Field Name"] = field
            row["Form Name"] = form
            row["Field Label"] = field_spec["label"]
            row["Required Field?"] = "y" if field_spec.get("required") else ""
            if kind == "category":
                row["Field Type"] = "radio"
                row["Choices, Calculations, OR Slider Labels"] = field_spec.get(
                    "choices", " | ".join(f"{value}, {value}" for value in field_spec["allowed"])
                )
            else:
                row["Field Type"] = "text"
            if kind == "int":
                row["Text Validation Type OR Show Slider Number"] = "integer"
                row["Text Validation Min"] = field_spec.get("min", "")
                row["Text Validation Max"] = field_spec.get("max", "")
            elif kind == "date":
                row["Text Validation Type OR Show Slider Number"] = "date_ymd"
                row["Text Validation Min"] = field_spec.get("min", "")
                row["Text Validation Max"] = field_spec.get("max", "")
                row["Identifier?"] = "y"
            if field_spec.get("pattern"):
                row["Field Annotation"] = f"Expected pattern: {field_spec['pattern']}"
            rows.append(row)
    return pd.DataFrame(rows, columns=DICTIONARY_COLUMNS)


def write_dictionary(rules, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    build_dictionary(rules).to_csv(path, index=False)
